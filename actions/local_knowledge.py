"""Explicit, local document indexing and retrieval for ARGUS."""

from __future__ import annotations

import os
import re
import sqlite3
import sys
import time
from html.parser import HTMLParser
from pathlib import Path
from typing import Iterable


SUPPORTED_SUFFIXES = {
    ".txt", ".md", ".rst", ".py", ".js", ".ts", ".tsx", ".jsx",
    ".json", ".jsonl", ".csv", ".tsv", ".html", ".htm", ".xml",
    ".yaml", ".yml", ".toml", ".ini", ".cfg", ".log", ".pdf",
    ".docx", ".pptx",
}
SKIP_DIRS = {
    ".git", ".hg", ".svn", ".venv", "venv", "env", "node_modules",
    "__pycache__", ".idea", ".vscode", ".ssh", "config", "configs",
    "credentials", "secrets", "private", "memory", "keys", "certs",
    "certificates",
}
SECRET_FILENAME = re.compile(
    r"(^\.env(?:\.|$)|credential|secret|private.?key|(^|_)id_rsa|"
    r"docker[-_.]?compose|(^|_)compose\.(?:ya?ml)$|token\.json)",
    re.IGNORECASE,
)
SECRET_CONTENT = re.compile(
    r"-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----|"
    r"\bAKIA[0-9A-Z]{16}\b|\bgh[pousr]_[A-Za-z0-9]{20,}\b|"
    r"\b(?:api[_-]?key|access[_-]?token|refresh[_-]?token|client[_-]?secret|"
    r"password|passwd|private[_-]?key|secret[_-]?key)\b"
    r"\s*[:=]\s*[\"']?[^\s\"']{8,}",
    re.IGNORECASE,
)
MAX_FILES_PER_INDEX = 2000
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_TOTAL_BYTES = 50 * 1024 * 1024
MAX_RESULTS = 5
MAX_SNIPPET_CHARS = 1100
CHUNK_MAX_CHARS = 4200
CHUNK_MAX_LINES = 80
CHUNK_OVERLAP_LINES = 6


class _TextExtractor(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        if data.strip():
            self.parts.append(data.strip())


def _base_dir() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys.executable).parent
    return Path(__file__).resolve().parent.parent


def _database_path(base_dir: Path | None = None) -> Path:
    return (base_dir or _base_dir()) / "memory" / "local_knowledge.sqlite3"


def _connect(db_path: Path) -> sqlite3.Connection:
    db_path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(str(db_path), timeout=15)
    connection.execute("PRAGMA busy_timeout = 15000")
    connection.execute("PRAGMA journal_mode = WAL")
    connection.executescript(
        """
        CREATE TABLE IF NOT EXISTS documents (
            path TEXT PRIMARY KEY,
            mtime_ns INTEGER NOT NULL,
            size_bytes INTEGER NOT NULL,
            indexed_at REAL NOT NULL,
            chunk_count INTEGER NOT NULL
        );
        CREATE VIRTUAL TABLE IF NOT EXISTS chunks USING fts5(
            content,
            path UNINDEXED,
            line_start UNINDEXED,
            line_end UNINDEXED,
            tokenize='unicode61 remove_diacritics 2'
        );
        """
    )
    return connection


def _is_sensitive_path(path: Path) -> bool:
    sensitive_dirs = {
        "config", "configs", "credentials", "secrets", "private", "memory",
        "keys", "certs", "certificates",
    }
    return any(
        part.startswith(".")
        or part.casefold() in sensitive_dirs
        or SECRET_FILENAME.search(part)
        for part in path.parts
    )


def _forget_document(connection: sqlite3.Connection, canonical_path: str) -> None:
    connection.execute("DELETE FROM chunks WHERE path = ?", (canonical_path,))
    connection.execute("DELETE FROM documents WHERE path = ?", (canonical_path,))


def _candidate_files(root: Path) -> Iterable[Path]:
    if root.is_file():
        yield root
        return
    for current, dirs, files in os.walk(root, topdown=True, followlinks=False):
        current_path = Path(current)
        dirs[:] = [
            name for name in dirs
            if not name.startswith(".")
            and name.casefold() not in SKIP_DIRS
            and not (current_path / name).is_symlink()
        ]
        for name in files:
            path = current_path / name
            if name.startswith(".") or path.is_symlink():
                continue
            yield path


def _extract_text(path: Path) -> str:
    suffix = path.suffix.casefold()
    if suffix in {".pdf", ".docx", ".pptx"}:
        if suffix == ".pdf":
            import fitz

            with fitz.open(path) as document:
                return "\n\n".join(
                    f"[Página {number}]\n{page.get_text()}"
                    for number, page in enumerate(document, 1)
                )
        if suffix == ".docx":
            from docx import Document

            document = Document(path)
            paragraphs = [item.text for item in document.paragraphs if item.text.strip()]
            for table in document.tables:
                paragraphs.extend(
                    " | ".join(cell.text.strip() for cell in row.cells)
                    for row in table.rows
                )
            return "\n".join(paragraphs)
        from pptx import Presentation

        presentation = Presentation(path)
        slides = []
        for number, slide in enumerate(presentation.slides, 1):
            text = "\n".join(
                shape.text for shape in slide.shapes
                if getattr(shape, "has_text_frame", False) and shape.text.strip()
            )
            slides.append(f"[Slide {number}]\n{text}")
        return "\n\n".join(slides)

    text = path.read_text(encoding="utf-8-sig", errors="replace")
    if suffix in {".html", ".htm"}:
        parser = _TextExtractor()
        parser.feed(text)
        return "\n".join(parser.parts)
    return text


def _split_chunks(text: str) -> list[tuple[int, int, str]]:
    lines = text.replace("\r\n", "\n").replace("\r", "\n").splitlines()
    chunks: list[tuple[int, int, str]] = []
    start = 0
    while start < len(lines):
        end = start
        size = 0
        while end < len(lines) and end - start < CHUNK_MAX_LINES:
            line_size = len(lines[end]) + 1
            if end > start and size + line_size > CHUNK_MAX_CHARS:
                break
            size += line_size
            end += 1
        content = "\n".join(lines[start:end]).strip()
        if content:
            chunks.append((start + 1, max(start + 1, end), content[:CHUNK_MAX_CHARS]))
        if end >= len(lines):
            break
        start = max(start + 1, end - CHUNK_OVERLAP_LINES)
    return chunks


def _resolve_requested_path(value: str) -> Path:
    if not value.strip():
        raise ValueError("Informe o caminho de um arquivo ou pasta para indexar.")
    path = Path(value).expanduser()
    if path.is_symlink():
        raise ValueError("ARGUS não indexa links simbólicos.")
    try:
        resolved = path.resolve(strict=True)
    except (OSError, RuntimeError) as exc:
        raise ValueError("O caminho informado não existe ou não pode ser acessado.") from exc
    home = Path.home().resolve()
    if resolved == home or home.is_relative_to(resolved):
        raise ValueError("Escolha uma pasta de documentos específica; ARGUS não indexa a pasta pessoal ou uma unidade inteira.")
    return resolved


def _index(path: Path, db_path: Path) -> str:
    if not path.exists() or not (path.is_file() or path.is_dir()):
        raise ValueError("O caminho precisa apontar para um arquivo ou pasta existente.")
    if _is_sensitive_path(path):
        raise ValueError("Esse caminho parece conter credenciais ou configuração sensível e não será indexado.")

    indexed = 0
    skipped = 0
    errors = 0
    total_bytes = 0
    seen = 0
    seen_paths: set[str] = set()
    scan_truncated = False
    connection = _connect(db_path)
    try:
        for file_path in _candidate_files(path):
            if seen >= MAX_FILES_PER_INDEX or total_bytes >= MAX_TOTAL_BYTES:
                skipped += 1
                scan_truncated = True
                break
            seen += 1
            try:
                canonical = str(file_path.resolve(strict=True))
                seen_paths.add(canonical)
            except (OSError, RuntimeError):
                errors += 1
                continue
            if file_path.suffix.casefold() not in SUPPORTED_SUFFIXES or _is_sensitive_path(file_path):
                _forget_document(connection, canonical)
                skipped += 1
                continue
            try:
                stat = file_path.stat()
                if stat.st_size > MAX_FILE_BYTES or total_bytes + stat.st_size > MAX_TOTAL_BYTES:
                    _forget_document(connection, canonical)
                    skipped += 1
                    continue
                total_bytes += stat.st_size
                text = _extract_text(file_path)
                if SECRET_CONTENT.search(text):
                    _forget_document(connection, canonical)
                    skipped += 1
                    continue
                chunks = _split_chunks(text)
                if not chunks:
                    _forget_document(connection, canonical)
                    skipped += 1
                    continue
                connection.execute("DELETE FROM chunks WHERE path = ?", (canonical,))
                connection.executemany(
                    "INSERT INTO chunks(content, path, line_start, line_end) VALUES (?, ?, ?, ?)",
                    [(content, canonical, line_start, line_end) for line_start, line_end, content in chunks],
                )
                connection.execute(
                    "INSERT INTO documents(path, mtime_ns, size_bytes, indexed_at, chunk_count) "
                    "VALUES (?, ?, ?, ?, ?) ON CONFLICT(path) DO UPDATE SET "
                    "mtime_ns=excluded.mtime_ns, size_bytes=excluded.size_bytes, "
                    "indexed_at=excluded.indexed_at, chunk_count=excluded.chunk_count",
                    (canonical, stat.st_mtime_ns, stat.st_size, time.time(), len(chunks)),
                )
                indexed += 1
            except Exception:
                _forget_document(connection, canonical)
                errors += 1
        if not scan_truncated:
            for (stored_path,) in connection.execute("SELECT path FROM documents").fetchall():
                stored = Path(stored_path)
                belongs_to_selection = stored == path or path in stored.parents
                if belongs_to_selection and stored_path not in seen_paths:
                    _forget_document(connection, stored_path)
        connection.commit()
    finally:
        connection.close()
    return (
        f"Indexação concluída: {indexed} arquivo(s) atualizado(s), "
        f"{skipped} ignorado(s) por tipo, limite ou proteção de dados, {errors} com erro. "
        "O conteúdo permanece no índice local do ARGUS."
    )


def _search(query: str, db_path: Path, top_k: int = MAX_RESULTS) -> str:
    terms = re.findall(r"[\wÀ-ÿ]{2,}", query, flags=re.UNICODE)[:12]
    if not terms:
        raise ValueError("Escreva uma busca com pelo menos uma palavra.")
    expression = " OR ".join(f'"{term}"' for term in terms)
    if not db_path.exists():
        return "Ainda não há documentos indexados. Peça para indexar uma pasta ou arquivo primeiro."
    connection = _connect(db_path)
    try:
        rows = connection.execute(
            "SELECT path, line_start, line_end, snippet(chunks, 0, '[', ']', ' … ', 18) "
            "FROM chunks WHERE chunks MATCH ? ORDER BY rank LIMIT ?",
            (expression, max(1, min(MAX_RESULTS, int(top_k)))),
        ).fetchall()
    except sqlite3.OperationalError as exc:
        if "fts5" in str(exc).lower() or "syntax" in str(exc).lower():
            raise ValueError("Não foi possível pesquisar esse texto no índice local.") from exc
        raise
    finally:
        connection.close()
    if not rows:
        return "Não encontrei trechos relevantes nos documentos indexados."
    results = ["Trechos encontrados nos documentos locais (use os caminhos como citações):"]
    for file_path, line_start, line_end, snippet in rows:
        excerpt = re.sub(r"\s+", " ", str(snippet or "")).strip()[:MAX_SNIPPET_CHARS]
        results.append(f"- {file_path}: linhas {line_start}–{line_end}\n  {excerpt}")
    return "\n".join(results)


def _list_documents(db_path: Path) -> str:
    if not db_path.exists():
        return "O índice local está vazio."
    connection = _connect(db_path)
    try:
        rows = connection.execute(
            "SELECT path, chunk_count FROM documents ORDER BY path LIMIT 100"
        ).fetchall()
        total = connection.execute("SELECT COUNT(*) FROM documents").fetchone()[0]
    finally:
        connection.close()
    if not rows:
        return "O índice local está vazio."
    output = [f"Documentos indexados ({total} no total):"]
    output.extend(f"- {path} ({count} trechos)" for path, count in rows)
    if total > len(rows):
        output.append(f"… e mais {total - len(rows)} arquivo(s).")
    return "\n".join(output)


def _remove(path_value: str, db_path: Path) -> str:
    if not path_value.strip():
        raise ValueError("Informe o caminho exato do arquivo ou pasta para remover do índice.")
    root = Path(path_value).expanduser()
    if root.is_symlink():
        raise ValueError("ARGUS não altera o índice por meio de links simbólicos.")
    root = root.resolve(strict=False)
    connection = _connect(db_path)
    try:
        rows = connection.execute("SELECT path FROM documents").fetchall()
        matches = []
        for (stored_path,) in rows:
            candidate = Path(stored_path)
            if candidate == root or root in candidate.parents:
                matches.append(stored_path)
        for stored_path in matches:
            connection.execute("DELETE FROM chunks WHERE path = ?", (stored_path,))
            connection.execute("DELETE FROM documents WHERE path = ?", (stored_path,))
        connection.commit()
    finally:
        connection.close()
    return f"Removi {len(matches)} documento(s) do índice local. Os arquivos originais não foram alterados."


def manage_local_knowledge(parameters: dict | None = None, *, base_dir: Path | None = None) -> str:
    """Index, search, list, or remove user-selected local documents."""
    args = parameters or {}
    action = str(args.get("action") or "search").strip().casefold()
    db_path = _database_path(base_dir)
    if action == "index":
        target = _resolve_requested_path(str(args.get("path") or ""))
        return _index(target, db_path)
    if action == "search":
        return _search(str(args.get("query") or ""), db_path, args.get("top_k", MAX_RESULTS))
    if action == "list":
        return _list_documents(db_path)
    if action == "remove":
        return _remove(str(args.get("path") or ""), db_path)
    raise ValueError("Ação inválida. Use index, search, list ou remove.")
